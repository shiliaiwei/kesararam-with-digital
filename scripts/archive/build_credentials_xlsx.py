import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

CRED = "/Applications/XAMPP/xamppfiles/htdocs/KWD/kesararam/scripts/generated_credentials.txt"
OUT = "/Applications/XAMPP/xamppfiles/htdocs/KWD/kesararam/Grade5_Student_Credentials_JVNUPL.xlsx"

rows = []
for line in open(CRED, encoding="utf-8"):
    p = [c.strip() for c in line.split("|")]
    if len(p) >= 6 and p[4] and p[4][0].isascii() and p[4][0].isalpha() and p[4] != "Username":
        no, khmer, fullname, gender, username, password = p[0], p[1], p[2], p[3], p[4], p[5]
        gmf = "M" if gender == "ប" else ("F" if gender == "ស" else gender)
        rows.append((khmer, fullname, gmf, username, password))

wb = Workbook()
ws = wb.active
ws.title = "Grade 5 - JVNUPL"

YELLOW = "EAB308"
BLACK = "000000"
GREY = "F4F4F5"
thin = Side(style="thin", color="000000")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

# Title
ws.merge_cells("A1:F1")
t = ws["A1"]
t.value = "Grade 5 Student Accounts  —  Class JVNUPL"
t.font = Font(name="Arial", bold=True, size=14, color=BLACK)
t.fill = PatternFill("solid", fgColor=YELLOW)
t.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 26

ws.merge_cells("A2:F2")
s = ws["A2"]
s.value = "Log in with Username + Password. Total: %d students." % len(rows)
s.font = Font(name="Arial", size=10, italic=True, color="3F3F46")
s.alignment = Alignment(horizontal="center")

headers = ["#", "Full name (English)", "Khmer name", "Gender", "Username", "Password"]
hrow = 3
for ci, h in enumerate(headers, start=1):
    c = ws.cell(row=hrow, column=ci, value=h)
    c.font = Font(name="Arial", bold=True, size=11, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=BLACK)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = border
ws.row_dimensions[hrow].height = 20

for i, (khmer, fullname, gmf, username, password) in enumerate(rows, start=1):
    r = hrow + i
    vals = [i, fullname, khmer, gmf, username, password]
    for ci, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=ci, value=v)
        c.border = border
        c.alignment = Alignment(horizontal="center" if ci in (1, 4) else "left", vertical="center")
        if ci == 5:
            c.font = Font(name="Consolas", size=11, color=BLACK)
        elif ci == 6:
            c.font = Font(name="Consolas", bold=True, size=11, color="B45309")
        else:
            c.font = Font(name="Arial", size=11, color=BLACK)
    if i % 2 == 0:
        for ci in range(1, 7):
            ws.cell(row=r, column=ci).fill = PatternFill("solid", fgColor=GREY)

widths = [5, 26, 24, 8, 24, 12]
for ci, w in enumerate(widths, start=1):
    ws.column_dimensions[chr(64 + ci)].width = w

ws.freeze_panes = "A4"
ws.sheet_view.showGridLines = False

wb.save(OUT)
print("Wrote", OUT, "with", len(rows), "students")
