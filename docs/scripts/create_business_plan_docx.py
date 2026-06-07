#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
KWD Business Plan DOCX Generator
Creates professional investor-ready business plan document
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

# Paths
TXT_FILE = '/Applications/XAMPP/xamppfiles/htdocs/KWD/kesararam/KWD_Business_Plan_5Y_Investor_Ready.txt'
LOGO_FILE = '/Applications/XAMPP/xamppfiles/htdocs/KWD/kesararam/public/brand/kwd-black_nobackground.png'
OUTPUT_FILE = '/Applications/XAMPP/xamppfiles/htdocs/KWD/kesararam/KWD_Business_Plan_5Y_Professional.docx'

def create_document():
    """Create and configure the document"""
    doc = Document()
    
    # Set document properties
    doc.core_properties.title = "KWD Platform - 5-Year Business Plan"
    doc.core_properties.author = "Kesararam With Digital"
    doc.core_properties.subject = "Investor-Ready Business Plan"
    doc.core_properties.keywords = "EdTech, Driving License, Cambodia, Business Plan"
    
    return doc

def add_header_with_logo(doc):
    """Add header with KWD logo"""
    section = doc.sections[0]
    header = section.header
    
    # Add logo to header
    if os.path.exists(LOGO_FILE):
        paragraph = header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(LOGO_FILE, width=Inches(2.0))
    
    return doc

def set_font(run, size=11, bold=False, khmer=False):
    """Set font properties"""
    font = run.font
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = RGBColor(0, 0, 0)
    
    # Use Khmer-compatible fonts
    if khmer:
        font.name = 'Khmer OS Battambang'  # Fallback for Khmer
    else:
        font.name = 'Arial'  # Fallback since Google Sans might not render properly
    
    return run

def parse_and_add_content(doc, txt_file):
    """Parse TXT file and add formatted content to document"""
    
    with open(txt_file, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    for line in lines:
        line = line.rstrip()
        
        # Skip decoration lines
        if line.startswith('════'):
            continue
        
        # Handle different line types
        if not line.strip():
            # Empty line
            doc.add_paragraph()
        
        elif line.startswith('│'):
            # Table content - add as formatted paragraph
            clean_line = line.replace('│', '').strip()
            if clean_line:
                p = doc.add_paragraph(clean_line)
                set_font(p.runs[0], size=10, khmer=True)
        
        elif line.startswith('┌') or line.startswith('├') or line.startswith('└'):
            # Box borders - skip or add as light separator
            doc.add_paragraph('_' * 80)
        
        elif line.startswith('១') or line.startswith('២') or line.startswith('៣'):
            # Main section headers (Khmer numbers)
            p = doc.add_heading(line, level=1)
            set_font(p.runs[0], size=16, bold=True, khmer=True)
        
        elif line.startswith('  •') or line.startswith('• ') or line.startswith('✓'):
            # Bullet points
            p = doc.add_paragraph(line.strip(), style='List Bullet')
            set_font(p.runs[0], size=11, khmer=True)
        
        elif '────' in line:
            # Subsection dividers - add as heading level 2
            clean_line = line.replace('────', '').strip()
            if clean_line and not clean_line.startswith('═'):
                p = doc.add_heading(clean_line, level=2)
                set_font(p.runs[0], size=14, bold=True, khmer=True)
        
        else:
            # Regular text
            p = doc.add_paragraph(line)
            if p.runs:
                set_font(p.runs[0], size=11, khmer=True)
    
    return doc

def add_page_numbers(doc):
    """Add page numbers to footer"""
    section = doc.sections[0]
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Add page number field
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    
    return doc

def main():
    """Main function to create business plan DOCX"""
    print("Creating KWD Business Plan DOCX...")
    
    # Create document
    doc = create_document()
    
    # Add header with logo
    doc = add_header_with_logo(doc)
    
    # Parse and add content from TXT file
    print("Parsing content from TXT file...")
    doc = parse_and_add_content(doc, TXT_FILE)
    
    # Add page numbers
    doc = add_page_numbers(doc)
    
    # Save document
    print(f"Saving document to: {OUTPUT_FILE}")
    doc.save(OUTPUT_FILE)
    
    print(f"✓ Business plan DOCX created successfully!")
    print(f"  File: {OUTPUT_FILE}")
    print(f"  Size: {os.path.getsize(OUTPUT_FILE) / 1024:.1f} KB")

if __name__ == "__main__":
    main()
